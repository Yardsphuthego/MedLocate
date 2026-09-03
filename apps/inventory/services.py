from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from .models import InventoryAdjustment, InventoryItem, InventoryUpload, InventoryUploadRow
from apps.audit.services import record_event
from apps.medicines.models import Medicine

REQUIRED_HEADERS = ("medicine_code", "medicine_name", "strength", "form", "quantity", "expiry_date")
OPTIONAL_HEADERS = ("last_updated",)


def _text(value):
    return " ".join(str(value or "").strip().split())


def _date_value(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value).strip())


def validate_inventory_workbook(*, uploaded_file, facility, actor):
    content = uploaded_file.read()
    digest = sha256(content).hexdigest()
    upload = InventoryUpload.objects.create(
        facility=facility, original_name=uploaded_file.name[:255], file_hash=digest, uploaded_by=actor
    )
    if InventoryUpload.objects.filter(facility=facility, file_hash=digest, status=InventoryUpload.Status.IMPORTED).exclude(pk=upload.pk).exists():
        upload.status = InventoryUpload.Status.REJECTED
        upload.error_text = "This exact workbook has already been imported for this facility."
        upload.save(update_fields=("status", "error_text"))
        return upload
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
        sheet = workbook.active
        rows = sheet.iter_rows()
        header_cells = next(rows)
    except (InvalidFileException, KeyError, ValueError, StopIteration, OSError) as exc:
        upload.status = InventoryUpload.Status.REJECTED
        upload.error_text = f"The workbook could not be read: {exc}"
        upload.save(update_fields=("status", "error_text"))
        return upload

    headers = tuple(_text(cell.value).lower() for cell in header_cells)
    missing = [name for name in REQUIRED_HEADERS if name not in headers]
    unexpected = [name for name in headers if name and name not in REQUIRED_HEADERS + OPTIONAL_HEADERS]
    if missing or unexpected or any(not name for name in headers) or len(set(headers)) != len(headers):
        parts = []
        if missing:
            parts.append("Missing headers: " + ", ".join(missing))
        if unexpected:
            parts.append("Unexpected headers: " + ", ".join(unexpected))
        if any(not name for name in headers):
            parts.append("Header names cannot be blank.")
        if len(set(headers)) != len(headers):
            parts.append("Header names must not be duplicated.")
        upload.status = InventoryUpload.Status.REJECTED
        upload.error_text = " ".join(parts)
        upload.save(update_fields=("status", "error_text"))
        return upload

    index = {name: headers.index(name) for name in headers if name}
    accepted, rejected, parsed, seen = 0, 0, [], set()
    for row_number, cells in enumerate(rows, start=2):
        if all(cell.value in (None, "") for cell in cells):
            continue
        errors = []
        if any(cell.data_type == "f" for cell in cells):
            errors.append("Formula cells are not allowed.")
        values = {name: cells[position].value if position < len(cells) else None for name, position in index.items()}
        code = _text(values.get("medicine_code"))
        name = _text(values.get("medicine_name"))
        strength = _text(values.get("strength"))
        medicine_form = _text(values.get("form"))
        try:
            raw_quantity = values.get("quantity")
            if isinstance(raw_quantity, bool):
                raise ValueError
            quantity = int(raw_quantity)
            if quantity < 0 or quantity != float(raw_quantity):
                raise ValueError
        except (TypeError, ValueError):
            quantity = None
            errors.append("Quantity must be a whole number of zero or more.")
        try:
            expiry = _date_value(values.get("expiry_date"))
        except (TypeError, ValueError):
            expiry = None
            errors.append("Expiry date must be a valid date.")
        if expiry and expiry < timezone.localdate() and quantity and quantity > 0:
            errors.append("Expired medicine cannot be imported as available stock.")
        medicine = Medicine.objects.filter(code=code, active=True).first()
        if not medicine:
            errors.append("Medicine code does not match an active medicine.")
        elif (_text(medicine.name).casefold(), _text(medicine.strength).casefold(), _text(medicine.form).casefold()) != (
            name.casefold(), strength.casefold(), medicine_form.casefold()
        ):
            errors.append("Medicine name, strength or form does not match the master record.")
        key = (code.casefold(), expiry.isoformat() if expiry else "")
        if key in seen:
            errors.append("Duplicate medicine and expiry date in this workbook.")
        seen.add(key)
        normalized = {
            "medicine_id": medicine.pk if medicine else None, "medicine_code": code, "medicine_name": name,
            "strength": strength, "form": medicine_form, "quantity": quantity,
            "expiry_date": expiry.isoformat() if expiry else None,
        }
        status = InventoryUploadRow.Status.REJECTED if errors else InventoryUploadRow.Status.ACCEPTED
        InventoryUploadRow.objects.create(
            upload=upload, row_number=row_number, status=status,
            error_text=" ".join(errors), normalized_data=normalized,
        )
        if errors:
            rejected += 1
        else:
            accepted += 1
            parsed.append(normalized)
    upload.accepted_count = accepted
    upload.rejected_count = rejected
    upload.parsed_rows = parsed
    upload.status = InventoryUpload.Status.VALIDATED if accepted else InventoryUpload.Status.REJECTED
    if not accepted and not upload.error_text:
        upload.error_text = "No valid inventory rows were found."
    upload.save(update_fields=("accepted_count", "rejected_count", "parsed_rows", "status", "error_text"))
    record_event(actor=actor, action="inventory.upload_validated", obj=upload, facility=facility,
                 metadata={"accepted": accepted, "rejected": rejected})
    return upload


@transaction.atomic
def commit_inventory_upload(*, upload, actor):
    upload = InventoryUpload.objects.select_for_update().get(pk=upload.pk)
    if upload.status != InventoryUpload.Status.VALIDATED:
        raise ValueError("Only a validated upload can be committed.")
    changed = 0
    for row in upload.parsed_rows:
        expiry = date.fromisoformat(row["expiry_date"]) if row["expiry_date"] else None
        item = InventoryItem.objects.select_for_update().filter(
            facility=upload.facility, medicine_id=row["medicine_id"], expiry_date=expiry
        ).first()
        previous = item.quantity if item else 0
        if item is None:
            item = InventoryItem(facility=upload.facility, medicine_id=row["medicine_id"], expiry_date=expiry)
        item.quantity = row["quantity"]
        item.updated_by = actor
        item.save()
        InventoryAdjustment.objects.create(
            item=item, previous_quantity=previous, new_quantity=item.quantity,
            reason=f"Snapshot import #{upload.pk}", actor=actor,
        )
        changed += 1
    upload.status = InventoryUpload.Status.IMPORTED
    upload.committed_at = timezone.now()
    upload.parsed_rows = []
    upload.save(update_fields=("status", "committed_at", "parsed_rows"))
    record_event(actor=actor, action="inventory.upload_committed", obj=upload, facility=upload.facility,
                 metadata={"records_changed": changed})
    return changed


@transaction.atomic
def save_inventory_item(*, form, facility, actor):
    item = form.save(commit=False)
    previous = 0
    if item.pk:
        previous = type(item).objects.select_for_update().get(pk=item.pk).quantity
    item.facility = facility
    item.updated_by = actor
    item.save()
    InventoryAdjustment.objects.create(
        item=item,
        previous_quantity=previous,
        new_quantity=item.quantity,
        reason=form.cleaned_data["reason"],
        actor=actor,
    )
    record_event(
        actor=actor,
        action="inventory.updated",
        obj=item,
        facility=facility,
        metadata={"previous_quantity": previous, "new_quantity": item.quantity},
    )
    return item
